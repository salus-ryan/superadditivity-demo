#!/usr/bin/env python3
"""Build a local, client-style sample prior-art evidence packet."""

from __future__ import annotations

import csv
import hashlib
import json
import os
import shutil
import subprocess
import textwrap
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

PROJECT = Path(__file__).resolve().parents[1]
JOB_ID = "patent-ocr-6fe76320-d10d-43a3-b6e1-4855780d9e6a"
BUNDLE = PROJECT / "jobs" / JOB_ID
INPUTS = BUNDLE / "inputs"
OUTPUT = PROJECT / "deliverables" / "sample-claim-evidence-packet"
ARCHIVE = PROJECT / "deliverables" / "sample-claim-evidence-packet.zip"

NAVY = colors.HexColor("#15253D")
BLUE = colors.HexColor("#2367A2")
PALE_BLUE = colors.HexColor("#EAF2F8")
TEAL = colors.HexColor("#147D7E")
PALE_TEAL = colors.HexColor("#E8F5F3")
AMBER = colors.HexColor("#B36B00")
PALE_AMBER = colors.HexColor("#FFF3D6")
RED = colors.HexColor("#A33A32")
PALE_RED = colors.HexColor("#FBEAE8")
GRAY = colors.HexColor("#5B6573")
LIGHT_GRAY = colors.HexColor("#EEF1F4")
DARK = colors.HexColor("#1E252C")

MATTER = {
    "id": "SAMPLE-SS-001",
    "title": "Preliminary Prior-Art Evidence Packet",
    "subtitle": "Synthetic distributed AI evidence-analysis claim",
    "search_date": "2026-10-14",
    "priority_cutoff": "2026-10-14",
    "disclaimer": (
        "Demonstration work product only. The target claim is fictional. This packet is a discovery and "
        "review aid, not legal advice, claim construction, a patentability opinion, or freedom-to-operate clearance."
    ),
}

ELEMENTS = [
    ("E1", "Partitioned multi-agent evidence", "Distinct evidence subsets are evaluated by corresponding AI evaluators."),
    ("E2", "Parallel typed probabilistic judgments", "Evaluators return typed judgments carrying probabilities or confidence values."),
    ("E3", "Connected verification", "A later stage compares, aggregates, or resolves outputs from multiple evaluators."),
    ("E4", "Deterministic policy composition", "A non-model policy mechanism maps judgments to a permitted action."),
    ("E5", "Coalitional contribution measurement", "The system measures marginal contribution, coalition value, or surplus."),
    ("E6", "Isolated-versus-connected visualization", "A graphical interface shows isolated and connected workflow states."),
    ("E7", "Tamper-evident minimized audit", "Minimized decision events are stored in a cryptographically verifiable audit record."),
]

TARGET_CLAIM = (
    "A computer-implemented evidence-analysis system comprising: (a) a plurality of isolated evidence silos, "
    "each providing a distinct evidence subset to a corresponding artificial-intelligence evaluator; (b) an "
    "evaluation service configured to execute the evaluators concurrently and receive from each evaluator a "
    "typed judgment and an associated probability; (c) a connected verifier configured to compare the typed "
    "judgments and produce a resolved determination; (d) a deterministic policy engine configured to map the "
    "resolved determination to a permitted response; (e) a utility engine configured to calculate a marginal "
    "contribution or coalition surplus associated with connecting the evidence silos; (f) a graphical interface "
    "configured to display isolated and connected workflow states; and (g) an audit service configured to record "
    "data-minimized decision events in a cryptographically chained audit log."
)

REFERENCES = [
    {
        "publication": "US12596813B2",
        "title": "Autonomous agent observation and control",
        "assignee": "Citibank, N.A.",
        "priority": "2023-01-19",
        "publication_date": "2026-04-07",
        "url": "https://patents.google.com/patent/US12596813B2/en",
        "profile": "Closest reviewed claim-language overlap for partitioned observation, connected comparison, conflict resolution, policy updates, and auditable records.",
        "coverage": "E1, E3, E4; partial E7",
        "image": ("US12596813B2", 66, "PDF page 66; publication pages 77–78"),
    },
    {
        "publication": "US12526244B2",
        "title": "Encrypted autonomous agent verification in multi-tiered distributed systems",
        "assignee": "Citibank, N.A.",
        "priority": "2024-02-06",
        "publication_date": "2026-01-13",
        "url": "https://patents.google.com/patent/US12526244B2/en",
        "profile": "Privacy-preserving multi-agent verification using transformed reference values, verification artifacts, records, and shared storage.",
        "coverage": "E1, E3",
        "image": ("US12526244B2", 62, "PDF page 62; publication pages 67–68"),
    },
    {
        "publication": "US20250259042A1",
        "title": "Platform for orchestrating a scalable, privacy-enabled network of collaborative and negotiating agents",
        "assignee": "Qomplx LLC",
        "priority": "2024-02-08",
        "publication_date": "2025-08-14",
        "url": "https://patents.google.com/patent/US20250259042A1/en",
        "profile": "Specialist agents process decomposed subtasks; a second plurality processes initial results into secondary results and a comprehensive response.",
        "coverage": "E1, E3",
        "image": ("US20250259042A1", 98, "PDF page 98; publication page 62"),
    },
    {
        "publication": "US20250328560A1",
        "title": "Security methods and systems for multi-agent generative AI applications",
        "assignee": "Vijay Madisetti",
        "priority": "2023-05-04",
        "publication_date": "2025-10-23",
        "url": "https://patents.google.com/patent/US20250328560A1/en",
        "profile": "Permission validation drives approve/block actions; a dependent claim recites an immutable registry and tamper-proof audit trails.",
        "coverage": "E4, E7",
        "image": ("US20250328560A1", 99, "PDF page 99; publication page 25"),
    },
    {
        "publication": "US20220210200A1",
        "title": "AI-driven defensive cybersecurity strategy analysis and recommendation system",
        "assignee": "Qomplx, Inc.",
        "priority": "2015-10-28",
        "publication_date": "2022-06-30",
        "url": "https://patents.google.com/patent/US20220210200A1/en",
        "profile": "A useful partial/negative control: simulated attack analysis produces and automatically implements recommendations, but reviewed claims do not recite the target multi-agent typed-judgment combination.",
        "coverage": "Partial E4 only",
        "image": ("US20220210200A1", 71, "PDF page 71; publication page 27"),
    },
]

EVIDENCE = [
    {
        "element": "E1", "publication": "US12596813B2", "claim": "Claim 1", "location": "PDF p. 66; pub. pp. 77–78",
        "quote": "receive, from a plurality of autonomous computing agents deployed within an operational environment, a plurality of activity data streams ... [and] deploy a plurality of observing agents",
        "assessment": "Express multi-agent/observer partitioning", "confidence": "High", "status": "Express",
    },
    {
        "element": "E1", "publication": "US20250259042A1", "claim": "Claim 1", "location": "PDF p. 98; pub. p. 62",
        "quote": "selecting appropriate artificial intelligence (AI) agents specializing in each of the domains ... [and] decomposing it into specialized subtasks pertaining to each of the selected AI agents",
        "assessment": "Express specialization and subtask partitioning", "confidence": "High", "status": "Express",
    },
    {
        "element": "E1", "publication": "US12526244B2", "claim": "Claim 1", "location": "PDF p. 62; pub. p. 68",
        "quote": "a first AI agent associated with a first entity ... a second AI agent associated with a second entity ... [and] a third AI agent associated with a third entity",
        "assessment": "Multiple agents with distinct entity relationships", "confidence": "High", "status": "Express",
    },
    {
        "element": "E2", "publication": "—", "claim": "—", "location": "Reviewed claim set",
        "quote": "No reviewed claim expressly required concurrent typed judgments carrying probability or confidence values.",
        "assessment": "Open gap in reviewed set; not an absence conclusion", "confidence": "Low", "status": "Not located",
    },
    {
        "element": "E3", "publication": "US12596813B2", "claim": "Claim 3", "location": "PDF p. 66; pub. p. 78",
        "quote": "compare ... the first observation summary ... with a second observation summary ... to detect a conflict ... [and] apply ... a conflict-resolution rule set comprising a precedence rule and a tie-breaking rule",
        "assessment": "Express connected comparison and resolution", "confidence": "High", "status": "Express",
    },
    {
        "element": "E3", "publication": "US20250259042A1", "claim": "Claim 1", "location": "PDF p. 98; pub. p. 62",
        "quote": "processing at least one plurality of AI agents' initial results through a second plurality of AI agents ... [and] developing a comprehensive response",
        "assessment": "Express second-stage processing/aggregation", "confidence": "High", "status": "Express",
    },
    {
        "element": "E4", "publication": "US20250328560A1", "claim": "Claim 1", "location": "PDF p. 98; pub. p. 24",
        "quote": "performing a permission validation ... and one of approving or blocking the requested operation ... based on the security analysis and the permission validation",
        "assessment": "Rule-like validation controls approve/block action", "confidence": "High", "status": "Express",
    },
    {
        "element": "E4", "publication": "US12596813B2", "claim": "Claim 3", "location": "PDF p. 66; pub. p. 78",
        "quote": "apply ... a conflict-resolution rule set comprising a precedence rule and a tie-breaking rule ... and update an observation policy",
        "assessment": "Explicit deterministic conflict-resolution rules", "confidence": "High", "status": "Express",
    },
    {
        "element": "E5", "publication": "—", "claim": "—", "location": "Reviewed claim set",
        "quote": "No reviewed claim expressly required coalition value, Shapley value, marginal contribution, or connected surplus measurement.",
        "assessment": "Open gap in reviewed set; not an absence conclusion", "confidence": "Low", "status": "Not located",
    },
    {
        "element": "E6", "publication": "—", "claim": "—", "location": "Reviewed claim set",
        "quote": "No reviewed claim expressly required an interface comparing isolated and connected workflow states.",
        "assessment": "Description-level dashboards were observed elsewhere; target claim limitation was not located", "confidence": "Low", "status": "Not located",
    },
    {
        "element": "E7", "publication": "US20250328560A1", "claim": "Claim 8", "location": "PDF p. 99; pub. p. 25",
        "quote": "storing a registry ... in an immutable blockchain registry ... and maintaining one or more tamper-proof audit trails of all security events",
        "assessment": "Express tamper-evident/immutable audit mechanism", "confidence": "High", "status": "Express",
    },
    {
        "element": "E7", "publication": "US12596813B2", "claim": "Claim 4", "location": "PDF p. 66; pub. p. 78",
        "quote": "store the mitigative action, a timestamp, and the decision path, on an auditable ledger",
        "assessment": "Audit record is express; tamper evidence is not express in this excerpt", "confidence": "Medium", "status": "Partial",
    },
]

REVIEW_QUEUE = [
    ("Critical", "Claim numbering", "Apple Vision omitted printed labels for claims 7 and 14 in US12596813B2.", "Verify all independent-claim boundaries against the authoritative record before client reliance.", "Open"),
    ("High", "E4 construction", "Permission validation and conflict-resolution rules may not satisfy every construction of deterministic policy composition.", "Attorney to define the target limitation and accept or reject each mapping.", "Open"),
    ("High", "E7 scope", "An 'auditable ledger' does not necessarily require cryptographic tamper evidence.", "Treat US12596813B2 as partial unless another claim or specification passage supplies the limitation.", "Open"),
    ("High", "Search completeness", "Only five documents received full OCR claim review; 46 publications were discovered.", "Expand CPC/IPC classes, families, and citation chains before any legal opinion.", "Open"),
    ("Medium", "Family grouping", "Current grouping is a title/priority proxy, not INPADOC family analysis.", "Validate family relationships in an authoritative database.", "Open"),
    ("Medium", "E2/E5/E6 gaps", "No express reviewed-claim match was located.", "Run targeted searches and inspect specifications; do not characterize these as absent from the art.", "Open"),
    ("Low", "Legal status", "Public metadata may be delayed or incomplete.", "Confirm current status and live claims by jurisdiction.", "Open"),
]

SEARCH_QUERIES = [
    "multi-agent evidence aggregation verification policy",
    "parallel probabilistic AI judgments confidence",
    "multi-agent deterministic policy approve block",
    "coalition utility Shapley multi-agent language model",
    "agent dashboard isolated connected workflow",
    "AI agent tamper-evident audit log",
    "autonomous agent conflict resolution auditable ledger",
]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def register_fonts() -> None:
    regular = Path("/system/fonts/Roboto-Regular.ttf")
    serif = Path("/system/fonts/NotoSerif-Regular.ttf")
    serif_bold = Path("/system/fonts/NotoSerif-Bold.ttf")
    if regular.exists():
        pdfmetrics.registerFont(TTFont("PacketSans", str(regular)))
    if serif.exists():
        pdfmetrics.registerFont(TTFont("PacketSerif", str(serif)))
    if serif_bold.exists():
        pdfmetrics.registerFont(TTFont("PacketSerifBold", str(serif_bold)))


def font(name: str, fallback: str) -> str:
    return name if name in pdfmetrics.getRegisteredFontNames() else fallback


def styles():
    base = getSampleStyleSheet()
    sans = font("PacketSans", "Helvetica")
    serif = font("PacketSerif", "Times-Roman")
    serif_bold = font("PacketSerifBold", "Times-Bold")
    return {
        "title": ParagraphStyle("TitleX", parent=base["Title"], fontName=serif_bold, fontSize=28, leading=33, textColor=NAVY, alignment=TA_LEFT, spaceAfter=14),
        "subtitle": ParagraphStyle("SubtitleX", parent=base["Normal"], fontName=sans, fontSize=13, leading=18, textColor=GRAY, spaceAfter=18),
        "h1": ParagraphStyle("H1X", parent=base["Heading1"], fontName=serif_bold, fontSize=20, leading=24, textColor=NAVY, spaceBefore=8, spaceAfter=10),
        "h2": ParagraphStyle("H2X", parent=base["Heading2"], fontName=serif_bold, fontSize=14, leading=18, textColor=BLUE, spaceBefore=8, spaceAfter=6),
        "body": ParagraphStyle("BodyX", parent=base["BodyText"], fontName=serif, fontSize=9.5, leading=14, textColor=DARK, spaceAfter=7),
        "small": ParagraphStyle("SmallX", parent=base["BodyText"], fontName=sans, fontSize=7.5, leading=10, textColor=GRAY),
        "callout": ParagraphStyle("CalloutX", parent=base["BodyText"], fontName=serif, fontSize=10, leading=15, textColor=NAVY, leftIndent=12, rightIndent=12, borderColor=BLUE, borderWidth=0.8, borderPadding=9, backColor=PALE_BLUE, spaceAfter=10),
        "quote": ParagraphStyle("QuoteX", parent=base["BodyText"], fontName=serif, fontSize=8.5, leading=12, textColor=DARK, leftIndent=9, borderColor=LIGHT_GRAY, borderWidth=0.5, borderPadding=7, backColor=colors.white),
        "covermeta": ParagraphStyle("CoverMeta", parent=base["BodyText"], fontName=sans, fontSize=9, leading=14, textColor=GRAY),
        "center": ParagraphStyle("CenterX", parent=base["BodyText"], fontName=sans, fontSize=9, leading=13, alignment=TA_CENTER, textColor=GRAY),
    }


def header_footer(canvas, doc) -> None:
    canvas.saveState()
    width, height = LETTER
    canvas.setStrokeColor(LIGHT_GRAY)
    canvas.line(0.65 * inch, height - 0.48 * inch, width - 0.65 * inch, height - 0.48 * inch)
    canvas.setFont(font("PacketSans", "Helvetica"), 7)
    canvas.setFillColor(GRAY)
    canvas.drawString(0.65 * inch, 0.38 * inch, f"{MATTER['id']} · DEMONSTRATION · NOT LEGAL ADVICE")
    canvas.drawRightString(width - 0.65 * inch, 0.38 * inch, f"Page {doc.page}")
    canvas.restoreState()


def status_color(status: str):
    return {"Express": PALE_TEAL, "Partial": PALE_AMBER, "Not located": PALE_RED}.get(status, colors.white)


def table(data, widths, *, header=True, font_size=7.2):
    t = Table(data, colWidths=widths, repeatRows=1 if header else 0, hAlign="LEFT")
    commands = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C8D0D8")),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("FONTNAME", (0, 0), (-1, -1), font("PacketSans", "Helvetica")),
        ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ("LEADING", (0, 0), (-1, -1), font_size + 2.5),
    ]
    if header:
        commands += [
            ("BACKGROUND", (0, 0), (-1, 0), NAVY),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), font("PacketSerifBold", "Times-Bold")),
        ]
    for row in range(1 if header else 0, len(data)):
        if row % 2 == 0:
            commands.append(("BACKGROUND", (0, row), (-1, row), colors.HexColor("#F7F9FA")))
    t.setStyle(TableStyle(commands))
    return t


def p(value: str, style) -> Paragraph:
    safe = value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return Paragraph(safe, style)


def render_evidence_pages() -> None:
    evidence_dir = OUTPUT / "evidence"
    evidence_dir.mkdir(parents=True, exist_ok=True)
    pages = {(ref["image"][0], ref["image"][1]) for ref in REFERENCES}
    pages.add(("US20250328560A1", 98))
    for publication, page in sorted(pages):
        destination = evidence_dir / f"{publication}-pdf-page-{page}.png"
        prefix = destination.with_suffix("")
        subprocess.run([
            "pdftoppm", "-f", str(page), "-l", str(page), "-singlefile",
            "-r", "100", "-png", str(INPUTS / f"{publication}.pdf"), str(prefix),
        ], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        os.chmod(destination, 0o600)


def build_executive_pdf(path: Path) -> None:
    s = styles()
    doc = SimpleDocTemplate(str(path), pagesize=LETTER, rightMargin=0.65*inch, leftMargin=0.65*inch, topMargin=0.68*inch, bottomMargin=0.62*inch, title=MATTER["title"], author="Silo Swarm demonstration")
    story = [Spacer(1, 0.55*inch), p("VERIFIED CLAIM REVIEW", s["covermeta"]), Spacer(1, 0.12*inch), p(MATTER["title"], s["title"]), p(MATTER["subtitle"], s["subtitle"])]
    story += [Table([
        [p("Matter", s["small"]), p(MATTER["id"], s["body"])],
        [p("Search date", s["small"]), p(MATTER["search_date"], s["body"])],
        [p("Corpus", s["small"]), p("46 discovered publications; five selected U.S. documents received full OCR claim review", s["body"])],
        [p("Target", s["small"]), p("One fictional independent claim decomposed into seven limitations", s["body"])],
        [p("Status", s["small"]), p("Preliminary discovery packet · human review required", s["body"])],
    ], colWidths=[1.25*inch, 5.45*inch], style=TableStyle([
        ("BOX", (0,0), (-1,-1), 0.8, LIGHT_GRAY), ("INNERGRID", (0,0), (-1,-1), 0.4, LIGHT_GRAY),
        ("BACKGROUND", (0,0), (0,-1), PALE_BLUE), ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("LEFTPADDING", (0,0), (-1,-1), 7), ("RIGHTPADDING", (0,0), (-1,-1), 7),
        ("TOPPADDING", (0,0), (-1,-1), 7), ("BOTTOMPADDING", (0,0), (-1,-1), 7),
    ])), Spacer(1, 0.25*inch), p(MATTER["disclaimer"], s["callout"]), Spacer(1, 0.2*inch), p("What the reviewer can answer in ten minutes", s["h2"])]
    for item in [
        "Which reviewed references contain the strongest claim-language overlap?",
        "Which target limitations were not expressly located in the reviewed claims?",
        "Where is every quoted passage in the authoritative source PDF?",
        "Which OCR, family, construction, and completeness issues still require professional review?",
    ]:
        story.append(p("• " + item, s["body"]))
    story += [PageBreak(), p("Executive finding", s["h1"]), p(
        "Two references dominate the reviewed claim-language overlap. US12596813B2 recites plural autonomous and observing agents, comparison of observation summaries, a conflict-resolution rule set, and an auditable ledger. US20250328560A1 recites permission validation controlling approval or blocking and, in dependent claim 8, an immutable blockchain registry with tamper-proof audit trails.", s["body"]), p(
        "US20250259042A1 separately recites domain-specialist agents, subtask decomposition, and second-stage processing of initial agent results. US12526244B2 recites privacy-preserving verification among multiple AI agents. No reviewed claim expressly required the target combination's typed probabilistic outputs, coalition-contribution measurement, or isolated-versus-connected visualization.", s["body"]), p(
        "The correct conclusion is therefore a prioritized review queue—not novelty. 'Not located' means only that the limitation was not expressly found in the five fully reviewed claim sets.", s["callout"]), p("At-a-glance limitation status", s["h2"])]
    summary_rows = [["Element", "Strongest reviewed evidence", "Status"]]
    summary_map = {
        "E1": ("US12596813B2 cl. 1; US20250259042A1 cl. 1", "Express"),
        "E2": ("No express reviewed-claim match", "Not located"),
        "E3": ("US12596813B2 cl. 3; US20250259042A1 cl. 1", "Express"),
        "E4": ("US20250328560A1 cl. 1; US12596813B2 cl. 3", "Express"),
        "E5": ("No express reviewed-claim match", "Not located"),
        "E6": ("No express reviewed-claim match", "Not located"),
        "E7": ("US20250328560A1 cl. 8", "Express"),
    }
    for key, name, _ in ELEMENTS:
        evidence, status = summary_map[key]
        summary_rows.append([f"{key} — {name}", evidence, status])
    t = table(summary_rows, [2.35*inch, 3.25*inch, 1.0*inch], font_size=7.5)
    for index, row in enumerate(summary_rows[1:], 1):
        t.setStyle(TableStyle([("BACKGROUND", (2,index), (2,index), status_color(row[2]))]))
    story += [t, PageBreak(), p("Fictional target claim", s["h1"]), p(
        "This synthetic claim exists only to demonstrate the review workflow. It is not recommended claim language and should not be published as an invention disclosure without authorization.", s["callout"]), p(TARGET_CLAIM, s["quote"]), Spacer(1, 0.15*inch), p("Claim decomposition", s["h2"])]
    for key, name, definition in ELEMENTS:
        story += [KeepTogether([p(f"{key} — {name}", s["h2"]), p(definition, s["body"])])]
    story += [PageBreak(), p("Element-by-element evidence chart", s["h1"])]
    for key, name, _ in ELEMENTS:
        story += [p(f"{key} — {name}", s["h2"])]
        rows = [["Reference / claim", "Quoted evidence", "Assessment"]]
        for item in [entry for entry in EVIDENCE if entry["element"] == key]:
            rows.append([
                p(f"{item['publication']}\n{item['claim']}\n{item['location']}", s["small"]),
                p(item["quote"], s["small"]),
                p(f"{item['status']} · {item['confidence']}\n{item['assessment']}", s["small"]),
            ])
        et = table(rows, [1.42*inch, 3.75*inch, 1.42*inch], font_size=6.8)
        for index, item in enumerate([entry for entry in EVIDENCE if entry["element"] == key], 1):
            et.setStyle(TableStyle([("BACKGROUND", (2,index), (2,index), status_color(item["status"]))]))
        story += [et, Spacer(1, 0.08*inch)]
    story += [PageBreak(), p("Reference ranking", s["h1"])]
    for index, ref in enumerate(REFERENCES, 1):
        story += [KeepTogether([
            p(f"{index}. {ref['publication']} — {ref['title']}", s["h2"]),
            p(f"Priority: {ref['priority']} · Assignee: {ref['assignee']} · Coverage: {ref['coverage']}", s["small"]),
            p(ref["profile"], s["body"]),
            p(f"Source: {ref['url']}", s["small"]),
        ])]
    story += [PageBreak(), p("Human review queue", s["h1"])]
    queue_rows = [["Priority", "Issue", "Why it matters", "Required action"]]
    for priority, issue, details, action, _ in REVIEW_QUEUE:
        queue_rows.append([priority, issue, p(details, s["small"]), p(action, s["small"])])
    qt = table(queue_rows, [0.62*inch, 1.05*inch, 2.35*inch, 2.58*inch], font_size=6.7)
    for row, item in enumerate(REVIEW_QUEUE, 1):
        fill = PALE_RED if item[0] == "Critical" else PALE_AMBER if item[0] == "High" else LIGHT_GRAY
        qt.setStyle(TableStyle([("BACKGROUND", (0,row), (0,row), fill)]))
    story += [qt, PageBreak(), p("Method and limits", s["h1"]), p(
        "Discovery combined structured query families, deduplication, public metadata, USPTO source PDFs, Apple Vision OCR, deterministic claim parsing, deterministic keyword charts, and a separate Jev semantic-triage pass. Source documents remained local during OCR. The semantic service received extracted claim text for technical disclosure triage.", s["body"]), p(
        "The workflow recorded source and OCR digests and used HMAC-chained controller events. These are SOC 2-aligned demo controls, not SOC 2 certification. The packet does not substitute for authoritative family analysis, prosecution-history review, jurisdiction-specific status checks, professional claim construction, or counsel's legal opinion.", s["body"]), p("Decision boundary", s["h2"]), p(
        "Use this packet to decide what a qualified reviewer should inspect next. Do not use automated coverage labels as an anticipation, obviousness, infringement, validity, or clearance conclusion.", s["callout"]), p("Verification", s["h2"]), p(
        "Run python3 verify.py from the packet directory. The verifier hashes every file listed in 05-Provenance-Manifest.json and reports any missing or modified artifact.", s["body"])]
    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)


def build_reference_pdf(path: Path) -> None:
    s = styles()
    doc = SimpleDocTemplate(str(path), pagesize=LETTER, rightMargin=0.6*inch, leftMargin=0.6*inch, topMargin=0.68*inch, bottomMargin=0.62*inch, title="Reference Profiles")
    story = [p("Reference profiles", s["title"]), p("Five selected publications with source-page evidence", s["subtitle"]), p(MATTER["disclaimer"], s["callout"]), PageBreak()]
    for index, ref in enumerate(REFERENCES):
        publication, page, citation = ref["image"]
        story += [p(f"{ref['publication']} — {ref['title']}", s["h1"]), table([
            ["Priority", ref["priority"], "Publication", ref["publication_date"]],
            ["Assignee", ref["assignee"], "Coverage", ref["coverage"]],
            ["Source", ref["url"], "Image", citation],
        ], [0.65*inch, 2.25*inch, 0.75*inch, 2.95*inch], header=False, font_size=7.2), Spacer(1, 0.1*inch), p(ref["profile"], s["body"])]
        image_path = OUTPUT / "evidence" / f"{publication}-pdf-page-{page}.png"
        from PIL import Image as PILImage
        with PILImage.open(image_path) as im:
            ratio = im.height / im.width
        width = 6.5 * inch
        height = min(6.55 * inch, width * ratio)
        story += [Image(str(image_path), width=width, height=height), p(f"Source image: {citation}. Full PDF is included under sources/.", s["center"])]
        if index < len(REFERENCES) - 1:
            story.append(PageBreak())
    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)


def build_protocol_pdf(path: Path) -> None:
    s = styles()
    doc = SimpleDocTemplate(str(path), pagesize=LETTER, rightMargin=0.7*inch, leftMargin=0.7*inch, topMargin=0.68*inch, bottomMargin=0.62*inch, title="Search Protocol")
    story = [p("Search and evidence protocol", s["title"]), p(f"Matter {MATTER['id']} · Search date {MATTER['search_date']}", s["subtitle"]), p(MATTER["disclaimer"], s["callout"]), p("Scope", s["h1"]), p(
        "The demonstration asks whether public patent claims located in a bounded discovery set expressly disclose seven elements of a fictional distributed AI evidence-analysis claim. Forty-six unique publications were retained from broader discovery; five selected U.S. documents received full OCR claim review.", s["body"]), p("Query families", s["h1"])]
    for query in SEARCH_QUERIES:
        story.append(p("• " + query, s["body"]))
    story += [p("Process", s["h1"])]
    for step in [
        "Acquire public metadata and preserve source URLs, identifiers, dates, titles, and snippets.",
        "Deduplicate publications and group proxy families by normalized title and priority date.",
        "Retrieve authoritative public PDF artifacts and bind each source with SHA-256.",
        "OCR image-only documents on the Mac worker; verify returned manifest and quality gates on the controller.",
        "Extract claim segments deterministically and preserve claim-text hashes.",
        "Map fixed elements using deterministic terms, then perform separate semantic technical-disclosure triage.",
        "Route missing claim labels, uncertain mappings, family questions, and legal construction to human review.",
    ]:
        story.append(p("• " + step, s["body"]))
    story += [p("Inclusion and grading", s["h1"]), table([
        ["Label", "Meaning"],
        ["Express", "The cited claim passage directly recites the mapped technical behavior."],
        ["Partial", "The passage overlaps but does not expressly supply the full target limitation."],
        ["Not located", "No express match was found in the reviewed claim set; this is not an absence conclusion."],
        ["Open", "The issue requires authoritative-source, family, status, or professional legal review."],
    ], [1.1*inch, 5.65*inch], font_size=8), p("Reproducibility", s["h1"]), p(
        "The packet includes all five source PDFs, rendered evidence pages, a source index, a machine-readable manifest, and a verifier. Raw OCR text is intentionally omitted from the client packet; quoted passages and page images are retained for review. Runtime logs remain local and are represented only through minimized counts and digests.", s["body"]), p("Limitations", s["h1"])]
    for limitation in [
        "One run and five full-text documents do not constitute an exhaustive patent search.",
        "OCR omitted some printed claim-number labels; authoritative claim boundaries require manual confirmation.",
        "The target claim is fictional and has not been legally construed.",
        "Automated semantic probabilities are triage signals, not legal findings.",
        "Public legal-status and family metadata may be incomplete or delayed.",
        "Unpublished applications and terminology outside the query families remain blind spots.",
    ]:
        story.append(p("• " + limitation, s["body"]))
    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)


def style_workbook(wb: Workbook) -> None:
    navy = "15253D"; white = "FFFFFF"; pale = "EAF2F8"; thin = Side(style="thin", color="CDD5DD")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(color=white, bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for index in range(1, ws.max_column + 1):
            max_len = max(len(str(ws.cell(row, index).value or "")) for row in range(1, min(ws.max_row, 50) + 1))
            ws.column_dimensions[get_column_letter(index)].width = min(55, max(11, max_len * 0.85))
        for row in range(2, ws.max_row + 1):
            if row % 2 == 0:
                for cell in ws[row]:
                    cell.fill = PatternFill("solid", fgColor=pale)
            ws.row_dimensions[row].height = 48
        ws.row_dimensions[1].height = 30


def build_claim_chart(path: Path) -> None:
    wb = Workbook()
    ws = wb.active; ws.title = "Element Chart"
    ws.append(["Element", "Target limitation", "Reference", "Claim", "Location", "Status", "Confidence", "Exact evidence", "Assessment", "Reviewer decision", "Reviewer notes"])
    element_map = {key: (name, definition) for key, name, definition in ELEMENTS}
    for item in EVIDENCE:
        name, definition = element_map[item["element"]]
        ws.append([f"{item['element']} — {name}", definition, item["publication"], item["claim"], item["location"], item["status"], item["confidence"], item["quote"], item["assessment"], "Pending", ""])
    refs = wb.create_sheet("Reference Index")
    refs.append(["Rank", "Publication", "Title", "Assignee", "Priority", "Publication date", "Coverage", "Profile", "Source URL"])
    for index, ref in enumerate(REFERENCES, 1): refs.append([index, ref["publication"], ref["title"], ref["assignee"], ref["priority"], ref["publication_date"], ref["coverage"], ref["profile"], ref["url"]])
    target = wb.create_sheet("Target Claim")
    target.append(["Matter", "Synthetic target claim", "Important notice"])
    target.append([MATTER["id"], TARGET_CLAIM, "Fictional demonstration claim; not recommended legal drafting."])
    queue = wb.create_sheet("Review Queue")
    queue.append(["Priority", "Issue", "Details", "Required action", "Status", "Reviewer", "Resolution"])
    for item in REVIEW_QUEUE: queue.append(list(item) + ["", ""])
    protocol = wb.create_sheet("Search Protocol")
    protocol.append(["Field", "Value"])
    protocol.append(["Search date", MATTER["search_date"]]); protocol.append(["Priority cutoff", MATTER["priority_cutoff"]]); protocol.append(["Corpus", "46 discovered publications; five selected documents received full OCR claim review"])
    for index, query in enumerate(SEARCH_QUERIES, 1): protocol.append([f"Query {index}", query])
    protocol.append(["Boundary", MATTER["disclaimer"]])
    style_workbook(wb)
    wb.save(path)


def build_review_queue(path: Path) -> None:
    wb = Workbook(); ws = wb.active; ws.title = "Review Queue"
    ws.append(["Priority", "Issue", "Details", "Required action", "Status", "Reviewer", "Reviewed at", "Resolution"])
    for item in REVIEW_QUEUE: ws.append(list(item) + ["", "", ""])
    instructions = wb.create_sheet("Instructions")
    instructions.append(["Step", "Instruction"])
    for index, instruction in enumerate([
        "Open the cited source PDF and rendered evidence page.",
        "Confirm claim number, page, and quotation against the authoritative record.",
        "Accept, modify, or reject each element mapping in 01-Claim-Chart.xlsx.",
        "Record legal construction and family/status corrections separately from machine extraction issues.",
        "Do not convert 'not located' into an absence, novelty, or clearance conclusion.",
    ], 1): instructions.append([index, instruction])
    style_workbook(wb); wb.save(path)


def copy_sources() -> None:
    sources = OUTPUT / "sources"; sources.mkdir(parents=True, exist_ok=True)
    for ref in REFERENCES:
        source = INPUTS / f"{ref['publication']}.pdf"
        destination = sources / source.name
        shutil.copy2(source, destination)
        os.chmod(destination, 0o600)
    with (OUTPUT / "source-index.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["publication", "priority", "publication_date", "assignee", "source_url", "local_file", "sha256"])
        for ref in REFERENCES:
            local = Path("sources") / f"{ref['publication']}.pdf"
            writer.writerow([ref["publication"], ref["priority"], ref["publication_date"], ref["assignee"], ref["url"], str(local), sha256_file(OUTPUT / local)])


def write_readme() -> None:
    (OUTPUT / "README.md").write_text(f"""# {MATTER['title']}

Matter: `{MATTER['id']}`

This local packet demonstrates a client-style prior-art evidence deliverable for one fictional claim. It is not legal advice, claim construction, a patentability opinion, or freedom-to-operate clearance.

## Start here

1. Read `00-Executive-Summary.pdf`.
2. Review and edit `01-Claim-Chart.xlsx`.
3. Open cited source pages through `02-Reference-Profiles.pdf` and `sources/`.
4. Read `06-Source-Verification-Record.md` and resolve remaining human-review issues in `04-Human-Review-Queue.xlsx`.
5. Run `python3 verify.py` to verify packet integrity.

## Important

- “Not located” means only not expressly found in the five fully reviewed claim sets.
- OCR omitted some printed claim-number labels; human verification remains mandatory.
- The target claim is synthetic and must not be published as an invention disclosure without authorization.
- Audit controls are SOC 2-aligned demo controls, not SOC 2 certification.
""", encoding="utf-8")


def write_source_verification() -> None:
    rows = []
    page_by_evidence = {
        ("US12596813B2", "Claim 1"): 66,
        ("US12596813B2", "Claim 3"): 66,
        ("US12596813B2", "Claim 4"): 66,
        ("US20250259042A1", "Claim 1"): 98,
        ("US12526244B2", "Claim 1"): 62,
        ("US20250328560A1", "Claim 1"): 98,
        ("US20250328560A1", "Claim 8"): 99,
    }
    for item in EVIDENCE:
        if item["publication"] == "—":
            continue
        page = page_by_evidence[(item["publication"], item["claim"])]
        image = OUTPUT / "evidence" / f"{item['publication']}-pdf-page-{page}.png"
        note = "Direct source image and normalized OCR phrase check passed."
        if item["publication"] == "US20250328560A1" and item["claim"] == "Claim 1":
            note = "Direct source image check passed; one normalized OCR segment crossed the page-end recognition boundary."
        rows.append(
            f"| {item['element']} | {item['publication']} | {item['claim']} | {item['location']} | PASS | `{sha256_file(image)}` | {note} |"
        )
    body = "\n".join(rows)
    (OUTPUT / "06-Source-Verification-Record.md").write_text(f"""# Source verification record

- Review date: 2026-10-14 UTC
- Reviewer: Pi coding agent, AI-assisted direct-source review
- Method: visual comparison against rendered USPTO PDF pages plus normalized phrase matching against Apple Vision OCR
- Result: 9/9 cited source rows located on the stated source pages
- Independent human/legal sign-off: pending

This closes the packet's machine/source traceability check. It does **not** represent attorney review, legal claim construction, or an independent human legal opinion.

| Element | Publication | Claim | Citation | Result | Evidence-image SHA-256 | Note |
|---|---|---|---|---|---|---|
{body}
""", encoding="utf-8")


def write_manifest() -> None:
    manifest_path = OUTPUT / "05-Provenance-Manifest.json"
    files = []
    for path in sorted(OUTPUT.rglob("*")):
        if not path.is_file() or path == manifest_path or path.name == "verify.py":
            continue
        files.append({"path": str(path.relative_to(OUTPUT)), "bytes": path.stat().st_size, "sha256": sha256_file(path)})
    manifest = {
        "schema_version": "1.0",
        "matter_id": MATTER["id"],
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z"),
        "generator": "scripts/build_sample_deliverable.py",
        "source_job_id": JOB_ID,
        "source_manifest_sha256": sha256_file(BUNDLE / "manifest.json"),
        "assurance_boundary": "SOC 2-aligned demo controls; not certification or an audit opinion.",
        "legal_boundary": MATTER["disclaimer"],
        "files": files,
    }
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    verifier = '''#!/usr/bin/env python3
import hashlib, json
from pathlib import Path
root = Path(__file__).resolve().parent
manifest = json.loads((root / "05-Provenance-Manifest.json").read_text())
failures = []
for item in manifest["files"]:
    path = root / item["path"]
    if not path.is_file():
        failures.append({"path": item["path"], "error": "missing"})
        continue
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if digest != item["sha256"]:
        failures.append({"path": item["path"], "error": "digest mismatch", "actual": digest})
print(json.dumps({"valid": not failures, "checked": len(manifest["files"]), "failures": failures}, indent=2))
raise SystemExit(1 if failures else 0)
'''
    (OUTPUT / "verify.py").write_text(verifier, encoding="utf-8")
    os.chmod(OUTPUT / "verify.py", 0o700)


def build_archive() -> None:
    ARCHIVE.unlink(missing_ok=True)
    with zipfile.ZipFile(ARCHIVE, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for path in sorted(OUTPUT.rglob("*")):
            if path.is_file():
                archive.write(path, Path(OUTPUT.name) / path.relative_to(OUTPUT))


def main() -> None:
    if not (BUNDLE / "manifest.json").is_file():
        raise SystemExit(f"missing source bundle: {BUNDLE}")
    shutil.rmtree(OUTPUT, ignore_errors=True)
    OUTPUT.mkdir(parents=True, mode=0o700)
    register_fonts()
    copy_sources()
    render_evidence_pages()
    build_executive_pdf(OUTPUT / "00-Executive-Summary.pdf")
    build_claim_chart(OUTPUT / "01-Claim-Chart.xlsx")
    build_reference_pdf(OUTPUT / "02-Reference-Profiles.pdf")
    build_protocol_pdf(OUTPUT / "03-Search-Protocol.pdf")
    build_review_queue(OUTPUT / "04-Human-Review-Queue.xlsx")
    write_readme()
    write_source_verification()
    write_manifest()
    build_archive()
    subprocess.run(["python3", str(OUTPUT / "verify.py")], check=True)
    print(json.dumps({
        "packet": str(OUTPUT),
        "archive": str(ARCHIVE),
        "archive_bytes": ARCHIVE.stat().st_size,
        "archive_sha256": sha256_file(ARCHIVE),
        "files": sum(1 for path in OUTPUT.rglob("*") if path.is_file()),
    }, indent=2))


if __name__ == "__main__":
    main()
